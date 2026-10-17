"""MS53-TC08/09/10/13: real ordinary-document parsing, complete support units."""
import hashlib

import pytest

from knowledge.chunk import chunk_sections
from knowledge.extract import extract
from knowledge.structure import manifest_hash, structure_hash
from test_knowledge_ingestion import CharacterTokenizer, text_pdf


ORDINARY = [
    'Welcome visitors at reception.\n\nAsk for their name and email before creating a booking.',
    'Opening routine\n\n1. Unlock the reception cabinet.\n2. Turn on the kiosk.\n3. Check the printed schedule.',
    'Lost property\n\n- Record the item description.\n- Store the item in the reception cabinet.\n- Never disclose another visitor\'s contact details.',
    'Can visitors bring food?\n\nFood is allowed in the pantry.\n\nCan visitors eat at desks?\n\nUse the pantry instead.',
    '| Item | Location |\n| --- | --- |\n| First aid | Reception |\n| Fire exit | East corridor |',
    'Closing the meeting room\n\nReturn borrowed adapters to reception.\n\nTurn off the display after everyone has left.',
    'Para sa receptionist\n\nKapag aalis ang customer, ibalik ang access card.\n\nHuwag magdagdag ng bayad nang walang pahintulot.',
    'Visitor Wi-Fi\n\nUse the network shown on the reception card.\n\nIf connection fails, contact the duty supervisor.',
]


def parsed_chunks(path, kind, scratch):
    result = extract(path, kind, scratch)
    chunks = chunk_sections(result['sections'], CharacterTokenizer())
    assert result['parser_fingerprint'].startswith('kb-parser-v5:')
    assert all(chunk['section_context']['structure_known'] for chunk in chunks)
    assert all(chunk['section_context']['scope'] == 'UNREVIEWED' for chunk in chunks)
    assert all(chunk['structure_sha256'] == structure_hash(chunk['section_context']) for chunk in chunks)
    assert len(manifest_hash(chunks)) == 64
    normalized = '\n\n'.join(section['text'] for section in result['sections'])
    for section in result['sections']:
        assert normalized[section['normalized_start']:section['normalized_end']] == section['text']
    assert result['content_sha256'] == hashlib.sha256(normalized.encode()).hexdigest()
    return result, chunks


@pytest.mark.parametrize('content', ORDINARY)
def test_ordinary_markdown_needs_no_headings_or_front_matter(tmp_path, content):
    path = tmp_path / 'ordinary.md'
    path.write_text(content, encoding='utf-8')
    _, chunks = parsed_chunks(path, 'MD', tmp_path)
    assert {chunk['section_context']['section_id'] for chunk in chunks} == {'s1'}
    assert all(not chunk['heading_path'] for chunk in chunks)
    assert '\n\n'.join(chunk['content'] for chunk in chunks) == content


@pytest.mark.parametrize('style', ['paragraphs', 'list', 'table'])
def test_real_docx_ordinary_styles_preserve_every_passage(tmp_path, style):
    from docx import Document
    document = Document()
    expected = ['Reception opens at eight.', 'Never leave the cash drawer unlocked.']
    if style == 'table':
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text, table.rows[0].cells[1].text = expected
    else:
        for text in expected:
            document.add_paragraph(text, style='List Bullet' if style == 'list' else None)
    path = tmp_path / 'ordinary.docx'
    document.save(path)
    _, chunks = parsed_chunks(path, 'DOCX', tmp_path)
    assert all(text in '\n'.join(chunk['content'] for chunk in chunks) for text in expected)
    assert {chunk['section_context']['section_id'] for chunk in chunks} == {'s1'}


@pytest.mark.parametrize('pages', [1, 2])
def test_real_text_pdf_keeps_page_provenance_and_whole_unit(tmp_path, pages):
    path = tmp_path / 'ordinary.pdf'
    # The helper writes real PDF objects, consumed by the approved pypdf parser.
    text_pdf(path)
    if pages == 2:
        from pypdf import PdfReader, PdfWriter
        reader = PdfReader(path)
        writer = PdfWriter()
        writer.add_page(reader.pages[0])
        writer.add_page(reader.pages[0])
        writer.write(path)
    result, chunks = parsed_chunks(path, 'PDF', tmp_path)
    assert {section['page'] for section in result['sections']} == set(range(1, pages + 1))
    assert {chunk['section_context']['section_id'] for chunk in chunks} == {'s1'}


def test_long_headingless_procedure_has_explicit_complete_support_unit(tmp_path):
    path = tmp_path / 'procedure.md'
    text = '\n\n'.join(f'{step}. Complete operation {step} without changing the next operation.' for step in range(1, 41))
    path.write_text(text)
    _, chunks = parsed_chunks(path, 'MD', tmp_path)
    assert len(chunks) > 1
    assert all(chunk['token_count'] <= 800 for chunk in chunks)
    assert {chunk['section_context']['section_id'] for chunk in chunks} == {'s1'}
    for step in range(1, 41):
        assert f'Complete operation {step} ' in '\n'.join(chunk['content'] for chunk in chunks)


def test_unresolved_references_still_require_correction(tmp_path):
    path = tmp_path / 'broken.md'
    path.write_text('Before doing this, see section Missing.\n\nDo the first step.')
    _, chunks = parsed_chunks(path, 'MD', tmp_path)
    assert chunks[0]['section_context']['unresolved_references']
    assert chunks[0]['section_context']['unresolved']
