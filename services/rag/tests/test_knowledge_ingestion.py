"""SPEC-5 synthetic extraction/security/worker tests. No operational documents."""
import hashlib
import json
from pathlib import Path
import subprocess
from types import SimpleNamespace
import zipfile

import pytest

from knowledge.chunk import chunk_sections, split_unit
from knowledge.extract import ExtractionError, extract, extract_docx, inspect_docx, normalize
from knowledge.worker import Gateway, GatewayError, process_claim


class CharacterTokenizer:
    def encode(self, text, add_special_tokens=True):
        return ([0] if add_special_tokens else []) + [ord(char) for char in text]
    def decode(self, values, **kwargs):
        return ''.join(chr(value) for value in values if value)


def test_spec5_tc02_markdown_preserves_sections_negation_and_unicode(tmp_path):
    source = tmp_path / 'source'
    source.write_text('# Access\n\nOnly staff may check in.\n\n## Reminder\nNever hold a desk.\n\nTagalog: hindi pa kumpirmado.', encoding='utf-8')
    result = extract(source, 'MD', tmp_path)
    assert result['sections'][0]['heading'] == 'Access'
    assert result['sections'][1]['heading'] == 'Access / Reminder'
    assert 'Never hold a desk.' in result['sections'][1]['text']
    assert result['source_sha256'] == hashlib.sha256(source.read_bytes()).hexdigest()
    assert not result['warning_codes']


def test_spec5_tc05_markdown_html_removed_links_not_fetched(tmp_path, monkeypatch):
    source = tmp_path / 'source'
    source.write_text('# SOP\n<script>fetch("https://private.invalid")</script>\n[Link](https://private.invalid)\n**Safe**')
    result = extract(source, 'MD', tmp_path)
    assert result['warning_codes'] == ['HTML_REMOVED']
    assert 'fetch(' not in result['sections'][0]['text']
    assert '[Link](https://private.invalid)' in result['sections'][0]['text']


@pytest.mark.parametrize('content', [b'', b'\xff\xfe\x00', b'MZ binary', b'PK\x03\x04binary'])
def test_spec5_tc05_invalid_markdown_rejected(tmp_path, content):
    source = tmp_path / 'source'
    source.write_bytes(content)
    with pytest.raises(ExtractionError):
        extract(source, 'MD', tmp_path)


def test_spec5_tc02_docx_order_headings_tables_and_footer(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_heading('Counter SOP', 1)
    doc.add_paragraph('Verify payment before allocation.')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'State'
    table.cell(0, 1).text = 'Action'
    table.cell(1, 0).text = 'Pending'
    table.cell(1, 1).text = 'Review'
    doc.add_paragraph('Finish the procedure.')
    doc.sections[0].footer.paragraphs[0].text = 'Internal guide'
    path = tmp_path / 'source.docx'
    doc.save(path)
    sections, warnings = extract_docx(path)
    joined = '\n'.join(section['text'] for section in sections)
    assert joined.index('Verify payment') < joined.index('| State | Action |') < joined.index('Finish')
    assert joined.count('Internal guide') == 1
    assert all(section['heading'] == 'Counter SOP' for section in sections[:-1])
    assert sections[-1]['heading'] == '' and sections[-1]['section_id'] == 's0'
    assert 'TABLE_REVIEW' in warnings


def unsafe_docx(path, extra_name, extra_bytes):
    from docx import Document
    Document().save(path)
    with zipfile.ZipFile(path, 'a') as archive:
        archive.writestr(extra_name, extra_bytes)


@pytest.mark.parametrize('name,content', [
    ('word/vbaProject.bin', b'not executed'),
    ('word/_rels/external.xml.rels', b'<Relationships><Relationship TargetMode="External" Target="https://private.invalid"/></Relationships>'),
    ('../outside', b'never extracted'),
    ('word/object.exe', b'MZ'),
    ('word/entity.xml', b'<!DOCTYPE test [<!ENTITY x SYSTEM "file:///etc/passwd">]><test/>'),
])
def test_spec5_tc05_unsafe_word_containers_fail(tmp_path, name, content):
    path = tmp_path / 'source.docx'
    unsafe_docx(path, name, content)
    with pytest.raises(ExtractionError, match='UNSAFE_FILE'):
        inspect_docx(path)


def test_spec5_tc04_embedded_word_object_requires_warning(tmp_path):
    path = tmp_path / 'source.docx'
    unsafe_docx(path, 'word/media/image.png', b'synthetic image')
    assert 'WORD_EMBEDDED_CONTENT' in inspect_docx(path)


def test_spec5_tc05_archive_budget_fails_instead_of_truncating(tmp_path, monkeypatch):
    from docx import Document
    path = tmp_path / 'source.docx'
    Document().save(path)
    monkeypatch.setenv('KB_ARCHIVE_EXPANDED_BYTES', '10')
    with pytest.raises(ExtractionError, match='RESOURCE_LIMIT'):
        inspect_docx(path)


def text_pdf(path, encrypted=False):
    from pypdf import PdfWriter
    from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(600, 800)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 12 Tf 50 700 Td (This is a synthetic operating procedure for text extraction.) Tj ET')
    page[NameObject('/Contents')] = stream
    if encrypted:
        writer.encrypt('synthetic-password')
    writer.write(path)


def test_spec5_tc02_pdf_text_page_provenance(tmp_path):
    source = tmp_path / 'source.pdf'
    text_pdf(source)
    result = extract(source, 'PDF', tmp_path)
    assert result['sections'][0]['page'] == 1
    assert result['sections'][0]['method'] == 'TEXT'
    assert 'synthetic operating procedure' in result['sections'][0]['text']


def test_spec5_tc05_encrypted_pdf_fails(tmp_path):
    source = tmp_path / 'source.pdf'
    text_pdf(source, encrypted=True)
    with pytest.raises(ExtractionError, match='ENCRYPTED_FILE'):
        extract(source, 'PDF', tmp_path)


def test_spec5_tc03_mixed_pdf_ocr_only_for_unreadable_page(tmp_path, monkeypatch):
    from pypdf import PdfReader, PdfWriter
    source = tmp_path / 'text.pdf'
    text_pdf(source)
    writer = PdfWriter()
    writer.add_page(PdfReader(source).pages[0])
    page = writer.add_blank_page(600, 800)
    from pypdf.generic import NameObject, DecodedStreamObject
    ink = DecodedStreamObject()
    ink.set_data(b'0 g 40 40 100 100 re f')
    page[NameObject('/Contents')] = writer._add_object(ink)
    mixed = tmp_path / 'mixed.pdf'
    writer.write(mixed)
    calls = []
    def ocr(arguments, **kwargs):
        if '--version' in arguments:
            return SimpleNamespace(returncode=0, stdout=b'tesseract synthetic-test-version', stderr=b'')
        calls.append(arguments)
        return SimpleNamespace(returncode=0, stdout=b'OCR synthetic Tagalog procedure: Suriin ang bayad.')
    monkeypatch.setattr(subprocess, 'run', ocr)
    result = extract(mixed, 'PDF', tmp_path)
    assert [section['method'] for section in result['sections']] == ['TEXT', 'OCR']
    assert len(calls) == 1 and calls[0][calls[0].index('-l') + 1] == 'eng+fil'
    assert 'OCR_REVIEW' in result['warning_codes']
    assert not list(tmp_path.glob('page-*.png'))


def test_spec5_tc16_unicode_long_text_not_truncated_and_hard_budget():
    text = ('予約 desk availability.\n' * 180) + 'END_OF_SOURCE'
    chunks = chunk_sections([{'heading': 'Guide', 'page': 3, 'method': 'OCR', 'text': text}], CharacterTokenizer())
    assert len(chunks) > 1
    assert all(0 < chunk['token_count'] <= 800 for chunk in chunks)
    assert all('\ufffd' not in chunk['content'] for chunk in chunks)
    assert chunks[-1]['content'].endswith('END_OF_SOURCE')
    assert [chunk['ordinal'] for chunk in chunks] == list(range(len(chunks)))
    assert all(chunk['page_start'] == 3 and chunk['extraction_method'] == 'OCR' for chunk in chunks)


def test_spec5_tc16_table_headers_repeat_and_role_context_survives():
    table = '| State | Action |\n| --- | --- |\n' + '\n'.join('| Pending | Review |' for _ in range(80))
    pieces = split_unit(table, CharacterTokenizer())
    assert len(pieces) > 1 and all(piece.startswith('| State | Action |\n| --- | --- |') for piece in pieces)
    chunks = chunk_sections([{'heading': 'Counter', 'page': None, 'method': 'TEXT', 'text': 'Only admins approve payment.\n\n' + table}], CharacterTokenizer())
    assert all('Only admins approve payment.' in chunk['content'] for chunk in chunks)


def test_spec5_tc15_worker_rejects_insecure_nonlocal_gateway():
    with pytest.raises(ValueError, match='PRIVATE_HTTPS'):
        Gateway('http://public.invalid', 'synthetic-worker-token-at-least-32-characters')


def test_spec5_tc08_lost_lease_never_completes(tmp_path):
    calls = []
    class LostGateway:
        def call(self, path, data):
            calls.append(path)
            raise GatewayError('LEASE_LOST')
    claim = {'job': {'id': 'job', 'kind': 'INGEST', 'generation': 1, 'lease_token': 'old'}, 'version': {}}
    process_claim(LostGateway(), claim, None, 'fingerprint')
    assert calls == ['job/checkpoint']


def test_spec5_tc14_purge_completes_only_through_gateway():
    calls = []
    class PurgeGateway:
        def call(self, path, data):
            calls.append((path, data))
    process_claim(PurgeGateway(), {'job': {'id': 'job', 'kind': 'PURGE', 'generation': 2, 'lease_token': 'current'}, 'version': {}}, None, 'fingerprint')
    assert calls == [('job/complete', {'generation': 2, 'lease_token': 'current'})]


def test_spec5_tc22_cli_never_prints_parser_exception_details(tmp_path):
    source = tmp_path / 'bad.docx'
    source.write_bytes(b'private synthetic source text')
    result = subprocess.run([str(Path(__import__('sys').executable)), '-m', 'knowledge.extract', '--source', str(source), '--type', 'DOCX'], capture_output=True, cwd=Path(__file__).resolve().parents[1], check=True)
    assert json.loads(result.stdout)['error']['code'] == 'MALFORMED_FILE'
    assert b'private synthetic' not in result.stdout + result.stderr
