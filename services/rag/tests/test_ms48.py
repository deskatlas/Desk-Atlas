"""MS48-TC01..08/13: real parsers and exact dependency-complete quotations."""
import hashlib
from types import SimpleNamespace

import pytest
from fastapi import HTTPException

from knowledge.chunk import chunk_sections
from knowledge.extract import extract
from knowledge.structure import manifest_hash, structure_hash
from knowledge.worker import process_claim
from test_assistant import api
from test_knowledge_ingestion import CharacterTokenizer, text_pdf


def large_sop():
    return '\n\n'.join([f'# Topic {i}\n\nOnly the Topic {i} supervisor approves Topic {i} changes.' for i in range(117)] + [
        '# Suggestions\n\nEmployees may suggest SOP improvements by submitting feedback to their supervisor.\n\nSee [Approval](#approval).',
        '# Approval\n\nOnly management may approve SOP changes. Suggestions do not authorize employees to modify the SOP.'])


def test_large_sop_preserves_unique_instances_and_separate_approval(tmp_path):
    path = tmp_path / 'source.md'
    path.write_text(large_sop())
    chunks = chunk_sections(extract(path, 'MD', tmp_path)['sections'], CharacterTokenizer())
    assert len(chunks) == 119
    assert len({c['section_context']['section_id'] for c in chunks}) == 119
    assert chunks[117]['section_context']['dependencies'] == ['s119']
    assert not any(c['section_context']['unresolved'] for c in chunks)
    assert all(c['structure_sha256'] == structure_hash(c['section_context']) for c in chunks)
    assert manifest_hash(chunks) != hashlib.sha256(''.join(c['content_sha256'] for c in chunks).encode()).hexdigest()


def test_repeated_headings_and_parent_conditions_are_not_merged(tmp_path):
    path = tmp_path / 'source.md'
    path.write_text('# Counter\nOnly admins approve.\n\n## Steps\nFirst procedure.\n\n## Steps\nSecond procedure.')
    chunks = chunk_sections(extract(path, 'MD', tmp_path)['sections'], CharacterTokenizer())
    assert len(chunks) == 3
    assert chunks[1]['heading_path'] == chunks[2]['heading_path']
    assert chunks[1]['section_context']['section_id'] != chunks[2]['section_context']['section_id']
    assert 's1' in chunks[2]['section_context']['dependencies']


def test_generic_policy_reference_does_not_infer_global_scope(tmp_path):
    path = tmp_path / 'source.md'
    path.write_text(large_sop().replace('Only the Topic 0 supervisor approves Topic 0 changes.',
                                      'Only the Topic 0 supervisor approves changes according to policy.'))
    chunks = chunk_sections(extract(path, 'MD', tmp_path)['sections'], CharacterTokenizer())
    assert len(chunks) == 119
    assert chunks[0]['section_context']['unresolved_references']
    assert chunks[0]['section_context']['scope'] == 'UNREVIEWED'
    assert all('s1' not in c['section_context']['dependencies'] for c in chunks[1:])
    assert chunks[117]['section_context']['dependencies'] == ['s119']


@pytest.mark.parametrize('reference', ['See section Missing.', 'See "Steps".', '[Missing](#absent)'])
def test_unknown_or_ambiguous_references_require_review(tmp_path, reference):
    path = tmp_path / 'source.md'
    path.write_text('# Query\n' + reference + '\n\n# Steps\nFirst.\n\n# Steps\nSecond.')
    chunks = chunk_sections(extract(path, 'MD', tmp_path)['sections'], CharacterTokenizer())
    assert chunks[0]['section_context']['unresolved']


def test_docx_and_pdf_structural_boundaries(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_heading('Procedure', 1)
    doc.add_paragraph('Step one.')
    doc.add_page_break()
    doc.add_paragraph('Step two.')
    doc.add_heading('Procedure', 1)
    doc.add_paragraph('Different procedure.')
    path = tmp_path / 'source.docx'
    doc.save(path)
    chunks = chunk_sections(extract(path, 'DOCX', tmp_path)['sections'], CharacterTokenizer())
    assert len(chunks) == 2
    assert 'Step two.' in chunks[0]['content']
    assert chunks[0]['section_context']['section_id'] != chunks[1]['section_context']['section_id']
    path = tmp_path / 'source.pdf'
    text_pdf(path)
    chunks = chunk_sections(extract(path, 'PDF', tmp_path)['sections'], CharacterTokenizer())
    # MS53-TC08: readable headingless PDFs now have a conservative whole-source unit.
    assert all(c['section_context']['structure_known'] for c in chunks)
    assert {c['section_context']['section_id'] for c in chunks} == {'s1'}


def test_selected_procedure_expands_conditions_but_not_unrelated_candidates():
    sources = [SimpleNamespace(chunk_id=str(i), content=text, heading_path='') for i, text in enumerate([
        'Employees may suggest improvements.', 'Only management may approve changes.', 'The kitchen opens at eight.'])]
    groups = [SimpleNamespace(id='suggest', chunk_ids=['0'], requires=['approve']),
              SimpleNamespace(id='approve', chunk_ids=['1'], requires=['suggest']),
              SimpleNamespace(id='kitchen', chunk_ids=['2'], requires=[])]
    body = SimpleNamespace(query='May employees suggest improvements?', sources=sources, groups=groups)
    result = api.complete_evidence_quotes(body, [{'chunk_id': '0', 'quote': sources[0].content}])
    assert [q['chunk_id'] for q in result] == ['0', '1']
    assert all(q['quote'] == sources[int(q['chunk_id'])].content for q in result)
    groups[0].requires.append('missing')
    with pytest.raises(HTTPException) as exc:
        api.complete_evidence_quotes(body, [{'chunk_id': '0', 'quote': sources[0].content}])
    assert exc.value.detail == 'UNRESOLVED_CONTEXT'


def test_complete_procedure_over_output_budget_is_not_no_source():
    body = SimpleNamespace(query='Procedure?', sources=[SimpleNamespace(chunk_id='0', content='x' * 1501, heading_path='')],
                           groups=[SimpleNamespace(id='procedure', chunk_ids=['0'], requires=[])])
    with pytest.raises(HTTPException) as exc:
        api.complete_evidence_quotes(body, [{'chunk_id': '0', 'quote': 'x'}])
    assert exc.value.detail == 'OUTPUT_BUDGET_EXCEEDED'


def test_markdown_code_headings_never_create_policy_sections(tmp_path):
    path = tmp_path / 'source.md'
    path.write_text('# Procedure\nActual text.\n\n```md\n# Example heading\n```\n\nContinuation.')
    chunks = chunk_sections(extract(path, 'MD', tmp_path)['sections'], CharacterTokenizer())
    assert len(chunks) == 1
    assert chunks[0]['heading_path'] == 'Procedure'


def test_long_structured_procedure_preserves_qualifier_without_repeated_tails(tmp_path):
    path = tmp_path / 'source.md'
    path.write_text('# Procedure\n\nOnly management approves changes.\n\n' + '\n\n'.join(f'Step {i}: inspect display number {i}.' for i in range(80)))
    chunks = chunk_sections(extract(path, 'MD', tmp_path)['sections'], CharacterTokenizer())
    assert len(chunks) > 1
    assert len({c['section_context']['section_id'] for c in chunks}) == 1
    assert sum('Only management approves changes.' in c['content'] for c in chunks) == 1
    for i in range(80):
        assert sum(f'Step {i}: inspect display number {i}.' in c['content'] for c in chunks) == 1


def test_context_selection_checks_identity_and_never_invents_chunk_ids(monkeypatch):
    import asyncio
    from unittest.mock import AsyncMock
    source = api.Source(document_id='doc', version_id='version', chunk_id='chunk', title='Synthetic',
                        heading_path='Procedure', source_page=None, epoch=1, extraction_method='TEXT',
                        content_hash='a' * 64, content='Employees may suggest improvements.')
    body = api.ContextInput(query='Suggest improvements?', sources=[source])
    monkeypatch.setattr(api, 'generate', AsyncMock(return_value=api.ContextIndexSelection(p=[3], a=False)))
    with pytest.raises(HTTPException) as exc:
        asyncio.run(api.context_selection(body, SimpleNamespace(ops=['answer'])))
    assert exc.value.detail == 'INVALID_MODEL_OUTPUT'
    with pytest.raises(HTTPException) as exc:
        asyncio.run(api.context_selection(body, SimpleNamespace(ops=['search'])))
    assert exc.value.status_code == 403


def test_context_selection_restores_four_ids_without_generating_uuid_tokens(monkeypatch):
    import asyncio
    import json
    from uuid import uuid4
    from unittest.mock import AsyncMock
    sources = [api.Source(document_id=str(uuid4()), version_id=str(uuid4()), chunk_id=str(uuid4()),
                          title='Synthetic checkout', heading_path=f'Step {i}', source_page=None, epoch=1,
                          extraction_method='TEXT', content_hash='a' * 64, content=f'Checkout step {i}.') for i in range(4)]
    async def select(_system, content, schema):
        wire = json.dumps(content)
        assert all(source.chunk_id not in wire and source.document_id not in wire for source in sources)
        assert [candidate['i'] for candidate in content['candidates']] == [0, 1, 2, 3]
        return schema.model_validate_json('{"p":[0,1,2,3],"a":false}')
    monkeypatch.setattr(api, 'generate', AsyncMock(side_effect=select))
    result = asyncio.run(api.context_selection(api.ContextInput(query='Checkout steps?', sources=sources), SimpleNamespace(ops=['answer'])))
    assert result == {'chunk_ids': [source.chunk_id for source in sources], 'ambiguous': False}


def test_empty_context_candidates_do_not_invoke_inference(monkeypatch):
    import asyncio
    from unittest.mock import AsyncMock
    generate = AsyncMock()
    monkeypatch.setattr(api, 'generate', generate)
    result = asyncio.run(api.context_selection(api.ContextInput(query='Checkout?', sources=[]), SimpleNamespace(ops=['answer'])))
    assert result == {'chunk_ids': [], 'ambiguous': False}
    generate.assert_not_called()


def test_old_parser_container_is_rejected_before_embedding(monkeypatch):
    import io
    import json
    from knowledge import worker
    raw = b'Synthetic source.'
    digest = hashlib.sha256(raw).hexdigest()
    calls = []
    class Gateway:
        def call(self, path, data):
            calls.append((path, data))
    class Encoder:
        @property
        def _model(self):
            pytest.fail('An incompatible parser must never reach embedding.')
    monkeypatch.setattr(worker, 'urlopen', lambda *_args, **_kwargs: io.BytesIO(raw))
    legacy = {'source_sha256': digest, 'content_sha256': digest,
              'parser_fingerprint': 'kb-parser-v1:synthetic', 'warning_codes': [],
              'sections': [{'heading': 'Legacy', 'page': None, 'method': 'TEXT', 'text': raw.decode()}]}
    monkeypatch.setattr(worker.subprocess, 'run', lambda *_args, **_kwargs: SimpleNamespace(returncode=0, stdout=json.dumps({'data': legacy}).encode()))
    claim = {'job': {'id': 'synthetic-job', 'kind': 'INGEST', 'generation': 1, 'lease_token': 'synthetic-lease'},
             'version': {'source_sha256': digest, 'source_bytes': len(raw), 'detected_type': 'MD'}, 'source_url': 'https://synthetic.invalid/source'}
    process_claim(Gateway(), claim, Encoder(), 'synthetic')
    assert calls[-1][1]['error_code'] == 'PARSER_REVISION_MISMATCH'
    assert not calls[-1][1]['retryable']
