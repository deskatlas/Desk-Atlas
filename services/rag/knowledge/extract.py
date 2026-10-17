"""SPEC5-TC02..05,22: isolated extraction. Never fetch links or execute source code."""
import argparse
import hashlib
import importlib.metadata
import json
import os
from pathlib import Path
import re
import subprocess
import tempfile
import unicodedata
import zipfile
from xml.etree import ElementTree


class ExtractionError(ValueError):
    pass


def digest(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def normalize(text: str) -> str:
    text = unicodedata.normalize("NFC", text).replace("\r\n", "\n").replace("\r", "\n")
    return "\n".join(line.rstrip() for line in text.splitlines()).strip()


def parser_fingerprint(kind: str) -> str:
    versions = {name: importlib.metadata.version(name) for name in ('python-docx', 'pypdf', 'pypdfium2', 'olefile')}
    versions['normalizer'] = 'nfc-v2-page-coverage'
    versions['chunker'] = 'kb-tokenizer-v5-ordinary-document-review'
    for name, command in (('libreoffice', ['libreoffice', '--version']), ('tesseract', ['tesseract', '--version'])):
        if (name == 'libreoffice' and kind == 'DOC') or (name == 'tesseract' and kind == 'PDF'):
            try:
                result = subprocess.run(command, capture_output=True, timeout=10, check=False)
                versions[name] = (result.stdout or result.stderr).decode('utf-8', errors='replace').splitlines()[0] if result.returncode == 0 else 'unavailable'
            except OSError:
                versions[name] = 'unavailable'
    if kind == 'PDF':
        data = Path(os.environ.get('TESSDATA_PREFIX', '/usr/share/tesseract-ocr/5/tessdata'))
        for language in ('eng', 'fil'):
            trained = data / (language + '.traineddata')
            versions[language] = digest(trained.read_bytes()) if trained.exists() else 'unavailable'
    return 'kb-parser-v5:' + digest(json.dumps(versions, sort_keys=True, separators=(',', ':')).encode())


def inspect_docx(path: Path) -> list[str]:
    warnings = set()
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            if not {"[Content_Types].xml", "word/document.xml"}.issubset(names):
                raise ExtractionError("TYPE_MISMATCH")
            # Engineering archive-bomb budget, not a product source-size cap.
            budget = int(os.environ.get("KB_ARCHIVE_EXPANDED_BYTES", str(256 * 1024 * 1024)))
            if sum(info.file_size for info in archive.infolist()) > budget:
                raise ExtractionError("RESOURCE_LIMIT")
            for info in archive.infolist():
                name = info.filename.lower()
                if info.flag_bits & 1:
                    raise ExtractionError("ENCRYPTED_FILE")
                if ".." in Path(name).parts or name.startswith("/") or len(names) != len(archive.infolist()):
                    raise ExtractionError("UNSAFE_FILE")
                if "vbaproject" in name or name.endswith((".exe", ".dll", ".js", ".vbs", ".ps1")):
                    raise ExtractionError("UNSAFE_FILE")
                if name.endswith(".rels"):
                    xml = archive.read(info)
                    if b"<!DOCTYPE" in xml.upper() or b"<!ENTITY" in xml.upper():
                        raise ExtractionError("UNSAFE_FILE")
                    for relation in ElementTree.fromstring(xml):
                        if relation.attrib.get("TargetMode") == "External":
                            raise ExtractionError("UNSAFE_FILE")
                if name.endswith(".xml"):
                    xml = archive.read(info)
                    if b"<!DOCTYPE" in xml.upper() or b"<!ENTITY" in xml.upper() or b"macroEnabled" in xml:
                        raise ExtractionError("UNSAFE_FILE")
                    if re.search(rb"<w:(ins|del|moveFrom|moveTo)\b", xml):
                        warnings.add("TRACKED_CHANGES")
                    if re.search(rb"<(w:txbxContent|w:object|w:altChunk|w:drawing)\b", xml):
                        warnings.add("WORD_EMBEDDED_CONTENT")
                if name.startswith(("word/media/", "word/embeddings/", "word/charts/")) or "comments" in name:
                    warnings.add("WORD_EMBEDDED_CONTENT")
    except (zipfile.BadZipFile, ElementTree.ParseError, RuntimeError, EOFError) as exc:
        raise ExtractionError("MALFORMED_FILE") from exc
    return sorted(warnings)


def extract_docx(path: Path) -> tuple[list[dict], list[str]]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    warnings = inspect_docx(path)
    document = Document(path)
    sections = []
    heading = []
    instances = []
    instance_count = 0
    counters = {}
    def append(text):
        if normalize(text):
            sections.append({"heading": " / ".join(heading), "page": None, "method": "TEXT", "text": normalize(text),
                             "section_id": instances[-1] if instances else 's0', "ancestors": ['s0'] + instances[:-1] if instances else []})
    for child in document.element.body:
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            style = paragraph.style.name if paragraph.style else ""
            text = paragraph.text
            if style.startswith("Heading ") and style.split()[-1].isdigit():
                level = int(style.split()[-1])
                instance_count += 1
                instances[:] = instances[:level - 1] + [f's{instance_count}']
                heading[:] = heading[:level - 1] + [text]
                append("#" * min(level, 6) + " " + text)
            else:
                numbering = paragraph._p.pPr.numPr if paragraph._p.pPr is not None else None
                if numbering is not None:
                    key = str(numbering.numId.val) if numbering.numId is not None else "default"
                    counters[key] = counters.get(key, 0) + 1
                    text = f"{counters[key]}. {text}"
                elif "List" in style:
                    text = "- " + text
                append(text)
        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            rows = [[cell.text.replace("\n", " / ").replace("|", "\\|") for cell in row.cells] for row in table.rows]
            if rows:
                append("\n".join(["| " + " | ".join(rows[0]) + " |", "| " + " | ".join("---" for _ in rows[0]) + " |"] + ["| " + " | ".join(row) + " |" for row in rows[1:]]))
                warnings.append("TABLE_REVIEW")
        elif not child.tag.endswith("}sectPr"):
            warnings.append("WORD_EMBEDDED_CONTENT")
    furniture = set()
    # Headers/footers have document scope, never the last procedure's scope.
    heading.clear()
    instances.clear()
    for section in document.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer, section.even_page_header, section.even_page_footer):
            for paragraph in part.paragraphs:
                if paragraph.text.strip() and paragraph.text not in furniture:
                    furniture.add(paragraph.text)
                    append(paragraph.text)
            if part.tables:
                warnings.append("WORD_EMBEDDED_CONTENT")
    return sections, sorted(set(warnings))


def extract_pdf(path: Path, scratch: Path) -> tuple[list[dict], list[str]]:
    from pypdf import PdfReader
    import pypdfium2 as pdfium
    reader = PdfReader(path)
    if reader.is_encrypted:
        raise ExtractionError("ENCRYPTED_FILE")
    if reader.trailer.get("/Root", {}).get("/OpenAction") or reader.trailer.get("/Root", {}).get("/AA"):
        raise ExtractionError("UNSAFE_FILE")
    renderer = None
    sections, warnings = [], set()
    try:
        for index, page in enumerate(reader.pages):
            text = normalize(page.extract_text() or "")
            method = "TEXT"
            readable = sum(character.isalnum() for character in text)
            # Images require coverage review even when a readable header exists.
            needs_ocr = bool(page.images) or readable < 20 or text.count("\ufffd") > max(2, len(text) // 100)
            if needs_ocr:
                if renderer is None:
                    renderer = pdfium.PdfDocument(str(path))
                rendered_page = renderer[index]
                bitmap = None
                image = scratch / f"page-{index}.png"
                try:
                    width, height = rendered_page.get_size()
                    if width <= 0 or height <= 0 or width * height * 4 > 24_000_000:
                        raise ExtractionError("CONTENT_LIMIT")
                    bitmap = rendered_page.render(scale=2)
                    pixels = bitmap.to_pil().convert("L")
                    # Empty OCR alone never proves blankness. Inspect every pixel.
                    if not text and pixels.getextrema()[0] >= 250:
                        sections.append({"heading": "", "page": index + 1, "method": "TEXT",
                                         "text": "", "outcome": "SKIPPED_BLANK"})
                        warnings.add("BLANK_PAGE_SKIPPED")
                        continue
                    pixels.save(image)
                    process = subprocess.run([os.environ.get("TESSERACT_PATH", "tesseract"), str(image), "stdout", "-l", "eng+fil", "--psm", "3"], capture_output=True, timeout=120, check=False)
                    if process.returncode:
                        raise ExtractionError("PARSER_FAILED")
                    ocr = normalize(process.stdout.decode("utf-8", errors="strict"))
                    if sum(char.isalnum() for char in ocr) < 3:
                        raise ExtractionError("UNREADABLE_PAGE")
                    # Keep native lines once; append OCR-only lines. Differences
                    # are visible for review instead of guessed corrections.
                    if text and readable >= 20:
                        native_lines = {re.sub(r'\s+', ' ', line).strip().casefold() for line in text.splitlines()}
                        additions = [line for line in ocr.splitlines()
                                     if re.sub(r'\s+', ' ', line).strip().casefold() not in native_lines]
                        text = normalize(text + '\n' + '\n'.join(dict.fromkeys(additions)))
                        method = "MIXED"
                        warnings.add("MIXED_CONTENT_REVIEW")
                    else:
                        text, method = ocr, "OCR"
                    warnings.add("OCR_REVIEW")
                finally:
                    image.unlink(missing_ok=True)
                    if bitmap is not None:
                        bitmap.close()
                    rendered_page.close()
            warnings.add("TABLE_REVIEW")
            sections.append({"heading": "", "page": index + 1, "method": method,
                             "text": normalize(text), "outcome": "EXTRACTED"})
    finally:
        if renderer is not None:
            renderer.close()
    return sections, sorted(warnings)


def extract(path: Path, kind: str, scratch: Path) -> dict:
    raw = path.read_bytes()
    warnings = []
    if kind == "MD":
        if raw.startswith((b"MZ", b"PK", b"%PDF-", b"\xd0\xcf\x11\xe0")) or b"\x00" in raw:
            raise ExtractionError("TYPE_MISMATCH")
        try:
            text = raw.decode("utf-8-sig", errors="strict")
        except UnicodeDecodeError as exc:
            raise ExtractionError("TYPE_MISMATCH") from exc
        # HTML is displayed as plain text in preview; remove raw blocks before indexing.
        text, removed = re.subn(r"<!--.*?-->|<(script|style)\b[^>]*>.*?</\1\s*>|<[^>]+>", "", text, flags=re.S | re.I)
        if removed:
            warnings.append("HTML_REMOVED")
        sections, heading, lines, instances = [], [], [], []
        instance_count = 0
        fence = None
        def flush():
            if normalize("\n".join(lines)):
                sections.append({"heading": " / ".join(heading), "page": None, "method": "TEXT", "text": normalize("\n".join(lines)),
                                 "section_id": instances[-1] if instances else 's0', "ancestors": ['s0'] + instances[:-1] if instances else []})
            lines.clear()
        for line in text.splitlines():
            marker = re.match(r'^\s{0,3}(`{3,}|~{3,})', line)
            if marker:
                if fence is None:
                    fence = marker[1][0]
                elif marker[1][0] == fence:
                    fence = None
                lines.append(line)
                continue
            if fence:
                lines.append(line)
                continue
            match = re.match(r"^(#{1,6})\s+(.+)$", line)
            if match:
                flush()
                instance_count += 1
                instances[:] = instances[:len(match[1]) - 1] + [f's{instance_count}']
                heading[:] = heading[:len(match[1]) - 1] + [match[2]]
            lines.append(line)
        flush()
    elif kind == "DOCX":
        sections, warnings = extract_docx(path)
    elif kind == "DOC":
        import olefile
        if not raw.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
            raise ExtractionError("TYPE_MISMATCH")
        with olefile.OleFileIO(str(path)) as office:
            streams = ["/".join(name).lower() for name in office.listdir()]
            if not office.exists("WordDocument"):
                raise ExtractionError("TYPE_MISMATCH")
            if any("vba" in name or "macros" in name for name in streams):
                raise ExtractionError("UNSAFE_FILE")
            word_header = office.openstream("WordDocument").read(12)
            if len(word_header) < 12 or int.from_bytes(word_header[10:12], "little") & 0x8100:
                raise ExtractionError("ENCRYPTED_FILE")
        original = scratch / "source.doc"
        original.write_bytes(raw)
        profile = scratch / "profile"
        (profile / "user").mkdir(parents=True)
        (profile / "user/registrymodifications.xcu").write_text('''<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop></item><item oor:path="/org.openoffice.Office.Writer/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>2</value></prop></item></oor:items>''', encoding="utf-8")
        converted = subprocess.run([os.environ.get("LIBREOFFICE_PATH", "libreoffice"), "-env:UserInstallation=" + profile.as_uri(), "--headless", "--norestore", "--convert-to", "docx:Office Open XML Text", "--outdir", str(scratch), str(original)], capture_output=True, timeout=120, check=False)
        if converted.returncode or not (scratch / "source.docx").exists():
            raise ExtractionError("PARSER_FAILED")
        sections, warnings = extract_docx(scratch / "source.docx")
        warnings.append("LEGACY_CONVERSION")
    elif kind == "PDF":
        if not raw.startswith(b"%PDF-"):
            raise ExtractionError("TYPE_MISMATCH")
        sections, warnings = extract_pdf(path, scratch)
    else:
        raise ExtractionError("TYPE_MISMATCH")
    if not sections or not any(section["text"].strip() for section in sections):
        raise ExtractionError("EMPTY_TEXT")
    # MS53-TC08/10: absent headings do not mean absent readable structure.
    # Keep the entire source as one support unit, with paragraph/list/table and
    # page boundaries retained. This conservative closure never guesses where
    # a procedure ends or treats a generated display label as a policy heading.
    # OCR/mixed reading order still requires focused verification.
    if not any(section.get('heading') for section in sections):
        for section in sections:
            if section['text'].strip():
                section['section_id'] = 's1'
                section['ancestors'] = []
    content = "\n\n".join(section["text"] for section in sections)
    offset = 0
    for section in sections:
        # Offsets address the content_sha256-bound normalized extraction, not
        # invented byte positions in binary Word/PDF source files.
        section['normalized_start'] = offset
        section['normalized_end'] = offset + len(section['text'])
        offset = section['normalized_end'] + 2
    return {"sections": sections, "warning_codes": sorted(set(warnings)), "source_sha256": digest(raw), "content_sha256": digest(content.encode()), "parser_fingerprint": parser_fingerprint(kind)}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--type", choices=["MD", "DOC", "DOCX", "PDF"], required=True)
    args = parser.parse_args()
    try:
        with tempfile.TemporaryDirectory(prefix="kb-") as directory:
            result = extract(args.source, args.type, Path(directory))
        print(json.dumps({"data": result}, ensure_ascii=False))
    except Exception as exc:
        code = str(exc) if isinstance(exc, ExtractionError) else "PARSER_FAILED"
        print(json.dumps({"error": {"code": code}}))
        # Never print parser exceptions, source filenames, or stderr.


if __name__ == "__main__":
    main()
